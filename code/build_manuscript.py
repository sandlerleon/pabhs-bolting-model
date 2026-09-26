# -*- coding: utf-8 -*-
"""Build the IJAMT manuscript. Every number is read from pabhs_results.json and
plc_verify_results.json; every DOI reference from the harvest_refs.py verification.

    python build_manuscript.py
"""
import io
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\PABHS_IJAMT"
R = json.load(io.open(os.path.join(HERE, "pabhs_results.json"), encoding="utf-8"))
V = json.load(io.open(os.path.join(HERE, "plc_verify_results.json"), encoding="utf-8"))
DOIREFS = json.load(io.open(os.path.join(HERE, "_refs.json"), encoding="utf-8"))
REPO_URL = "https://github.com/sandlerleon/pabhs-bolting-model"
CODE_DOI = os.environ.get("PABHS_CODE_DOI")      # set once the Zenodo deposit exists

TITLE = ("Automated bolting and head handling for multi-flange pressure-vessel closures: "
         "a simulation-based design evaluation with a delayed-coking case study")

STANDARDS = {
    "pcc1": "ASME (2022) PCC-1-2022 Guidelines for pressure boundary bolted flange joint "
            "assembly. American Society of Mechanical Engineers, New York",
    "iso898": "ISO (2013) ISO 898-1:2013 Mechanical properties of fasteners made of carbon "
              "steel and alloy steel. Part 1: Bolts, screws and studs with specified property "
              "classes. International Organization for Standardization, Geneva",
    "iso261": "ISO (1998) ISO 261:1998 ISO general purpose metric screw threads. General "
              "plan. International Organization for Standardization, Geneva",
    "astm193": "ASTM International (2023) ASTM A193/A193M-23 Standard specification for "
               "alloy-steel and stainless steel bolting for high temperature or high pressure "
               "service and other special purpose applications. ASTM International, West "
               "Conshohocken",
    "iso15552": "ISO (2018) ISO 15552:2018 Pneumatic fluid power. Cylinders with detachable "
                "mountings, 1 000 kPa (10 bar) series, bores from 32 mm to 320 mm. Basic, "
                "mounting and accessories dimensions. International Organization for "
                "Standardization, Geneva",
    "iec60079": "IEC (2017) IEC 60079-0:2017 Explosive atmospheres. Part 0: Equipment. "
                "General requirements. International Electrotechnical Commission, Geneva",
    "atex": "European Parliament and Council (2014) Directive 2014/34/EU on the "
            "harmonisation of the laws of the Member States relating to equipment and "
            "protective systems intended for use in potentially explosive atmospheres. "
            "Official Journal of the European Union L 96:309-356",
    "iec61511": "IEC (2016) IEC 61511-1:2016 Functional safety. Safety instrumented systems "
                "for the process industry sector. Part 1: Framework, definitions, system, "
                "hardware and application programming requirements. International "
                "Electrotechnical Commission, Geneva",
}

# ------------------------------------------------------------------ document
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.5
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1.0)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
    s._sectPr.find(qn("w:pgMar")).addnext(ln)

ORDER = []


def C(*tags):
    nums = []
    for t in tags:
        if t not in ORDER:
            ORDER.append(t)
        nums.append(ORDER.index(t) + 1)
    nums = sorted(set(nums))
    # compress runs: [3,4,5] -> 3-5
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(str(nums[i]) if j - i < 2 else "%d\u2013%d" % (nums[i], nums[j]))
        if j - i == 1:
            parts[-1] = str(nums[i]); parts.append(str(nums[j]))
        i = j + 1
    return "[%s]" % ", ".join(parts)


def H(text, size=12, before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)


def H2(text):
    H(text, size=11, before=8)


def Pp(text, indent=True, italic=False, size=11, spacing=1.5, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = italic; r.font.size = Pt(size)
    return p


EQN = [0]


def EQ(text):
    EQN[0] += 1
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Inches(0.6)
    r = p.add_run(text); r.italic = True
    p.add_run("\t\t(%d)" % EQN[0])
    return EQN[0]


def CAP(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text); r.font.size = Pt(9.5)


def FIG(name, width=6.3):
    doc.add_picture(os.path.join(HERE, "figures", name + ".png"), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def TBL(headers, rows, fs=8.5, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(fs)
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_after = Pt(0)
            pp.add_run(str(v)).font.size = Pt(fs)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Inches(w)


cyc = R["cycle"]
b1 = cyc["1 head, parallel flanges"]
b2 = cyc["2 heads, parallel flanges"]
b4 = cyc["4 heads, parallel flanges"]
b4r = cyc["4 heads, 2 star passes"]
bs = cyc["1 head, serial flanges"]
bc = cyc["1 head, circular passes only"]
best = R["best_case_bolting_h"]
adm = R["admissibility"]
sens1 = R["sensitivity_bolting_total"]
sens4 = R["sensitivity_bolting_total_4heads"]
man = {row["head_t"]: row for row in R["manipulator"]}
fm = R["fmea"]
ex = R["exposure"]
pl = R["preload"]
tq = pl["torque_kNm"]
ta = {(t["cv_K"], t["cv_theta"]): t for t in pl["torque_angle"]}
grid = {(g["cv_K"], g["tool_pm"]): g for g in pl["torque_control_grid"]}
fd12 = V["full_design"]["12"]; as12 = V["as_specified"]["12"]
fd4 = V["full_design"]["4"]; as4 = V["as_specified"]["4"]
OPT = R["optimisation"]
OA, OB = OPT["A 4 heads, single carousel"], OPT["B + 4 transfer channels"]
OC, OD, OE = OPT["C + 2 star passes"], OPT["D + optimised sequence"], OPT["E + pipelined servicing"]
O1 = OPT["1 head, optimised sequence"]
COMP = OPT["E_composition_h"]
SQ = R["sequence_optimisation"]


def mr(x):
    return "%.2f h (90%% interval %.2f\u2013%.2f h)" % (x["median"], x["p05"], x["p95"])


# ==================================================================== FRONT
p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.3
r = p.add_run(TITLE); r.bold = True; r.font.size = Pt(14)
Pp("Leon Sandler", indent=False, spacing=1.2)
Pp("Independent Researcher, Northbrook, Illinois, USA", indent=False, size=10, spacing=1.2)
Pp("E-mail: sandler.leon@gmail.com \u00b7 ORCID: 0009-0007-4584-808X", indent=False, size=10,
   spacing=1.2)

H("Abstract")
ABSTRACT = (
    "Large bolted closures on process vessels are still assembled largely by "
    "hand, often in hazardous areas, using multi-pass tightening sequences whose duration "
    "and preload outcome are rarely quantified. This paper evaluates by simulation "
    "an automated architecture for such closures: orbital pneumatic torque runners "
    "on the flange, a bolt-management carousel, a counterbalanced head manipulator and an "
    "interlocking sequence controller, applied to a 92-bolt, three-flange delayed-coking drum. "
    "A Monte Carlo model built on the actual star-pattern tool travel shows that one torque "
    "head per flange needs a median %.1f h to unbolt and re-bolt the drum and, within the "
    "investigated parameter bounds, cannot meet a 2.5 h target. Removing bottlenecks in turn "
    "gives %.1f h with four synchronised heads and %.1f h with a bolt-transfer channel per "
    "head, the recommended machine-only configuration, which meets 2.5 h in %.0f%% of "
    "samples; one fewer star pass would give %.1f h but requires qualification. This ranking "
    "holds in every sample under uniform, triangular and correlated input models. An exactly "
    "optimised tool path adds little, and the complete operation remains about %.1f h, "
    "governed by head handling and leak testing. Specifying \u00b12%% torque accuracy barely "
    "affects preload scatter; under an assumed 3%% angle-based preload-estimate uncertainty, "
    "torque\u2013angle re-torque cuts bolts outside \u00b110%% of target from %.0f%% to "
    "%.0f%%. Exhaustive verification found no safety violation but exposed deadlocks after "
    "interrupted sequences, which a resume rule removed. All results are predictions; an "
    "experimental falsification plan is given."
    % (b1["bolting_total"]["median"], OA["bolting_total"]["median"],
       OB["bolting_total"]["median"], 100 * OB["p_bolting_le_2.5h"],
       OC["bolting_total"]["median"], OE["full_total"]["median"],
       100 * ta[(0.1, 0.03)]["outside10_before"], 100 * ta[(0.1, 0.03)]["outside10_after"]))
Pp(ABSTRACT, indent=False)
Pp("Keywords: automated bolting; bolted flange joint; torque\u2013angle control; Monte Carlo "
   "simulation; interlock verification; hazardous-area automation", indent=False, size=10)

# ==================================================================== 1
H("1 Introduction")
Pp("Bolted flange closures are the most frequently disturbed pressure boundaries in process "
   "plants, and their assembly is a manufacturing operation in all but name: a sequence of "
   "positioning, fastening and verification steps whose quality depends on the order in which "
   "fasteners are loaded and on how the applied torque converts to preload. Guidance for "
   "pressure-boundary joints prescribes multi-pass tightening sequences for exactly this "
   "reason %s, because tightening one bolt changes the load in its neighbours through elastic "
   "interaction %s. Where such joints are opened every few hours, the assembly procedure "
   "becomes a recurring production task, and the case for automating it is at once a quality, "
   "a throughput and a safety question." % (C("pcc1"), C("bibel1992", "zhu2018")))
Pp("Delayed-coking drums are an extreme instance. Each drum is a large pressure vessel "
   "subject to severe thermal cycling and a range of damage mechanisms %s, and its heads are "
   "unbolted and re-bolted on every coking cycle unless automatic unheading devices have been "
   "fitted. Bottom-head flange leaks are a common problem, the industry trend is to install "
   "automatic bottom unheading devices, and the joints that remain require periodic "
   "re-tightening and careful assembly practice %s. Thermal transients can also distort the "
   "drum flange until bolt holes no longer align %s. In the unit considered here, four drums "
   "on a 48 h fill cycle give an opening or closing operation every 12 h, and each operation "
   "involves 92 bolts across three flanged connections, handled manually at height in a "
   "potentially flammable, hot and dusty atmosphere." % (C("ma2023"), C("berry2019"),
                                                         C("horstmeyer2025")))
Pp("An automated architecture for this operation was first proposed by the author in "
   "response to an industrial open-innovation challenge, as a pneumatic and automated bolt "
   "handling system (PABHS) with four subsystems: orbital torque runners, a bolt carousel, a "
   "head manipulator and a sequence controller. That proposal stated performance targets "
   "\u2014 bolting in 1.5\u20132.5 h against a reported 4\u20136 h manually, \u00b12% torque "
   "accuracy and no personnel in the hazard zone \u2014 without the models needed to "
   "establish whether the architecture could meet them. This paper supplies those models. It "
   "treats the proposal as a design hypothesis and asks, for each target, what the "
   "architecture would have to achieve and under what conditions it would fail.")
Pp("The contributions are: (i) a cycle-time model for automated multi-pass flange bolting "
   "that computes carriage travel from the actual tightening sequence and propagates step-time "
   "uncertainty by Monte Carlo simulation, with a rank-correlation sensitivity analysis; (ii) "
   "an admissibility analysis giving the per-bolt-pass time budget implied by a cycle-time "
   "target; (iii) a throughput-optimisation study that removes the architecture's bottlenecks "
   "one at a time \u2014 torque heads, bolt-transfer channels, tightening passes, an exactly "
   "optimised tool path under a load-spreading constraint, and pipelined servicing \u2014 and "
   "reports bolting and complete-operation times separately, with a robustness check against "
   "the assumed input distributions, a design feasibility map and a servicing-time envelope "
   "for the complete operation; (iv) a preload analysis "
   "separating the contributions of tool accuracy, nut-factor scatter and torque\u2013angle "
   "verification; (v) exhaustive verification of the sequence and interlock logic, including "
   "mutation tests and a liveness check that exposes a deadlock in the logic as first "
   "specified; and (vi) a force budget for pneumatic head handling, a scored design FMEA and "
   "a hazard-exposure comparison. All models are released as open code, and the paper ends "
   "with a validation plan whose acceptance criteria can falsify its predictions.")

# ==================================================================== 2
H("2 Related work")
H2("2.1 Flange assembly and tightening sequences")
Pp("The legacy tightening pattern for pressure-boundary flanges loads bolts in a cross "
   "(star) order over several passes at increasing torque, followed by circular check passes "
   "until nuts no longer rotate %s. Alternative patterns that reduce assembly time or improve "
   "preload uniformity were added to that guidance with supporting laboratory evidence %s, "
   "and experimental comparisons of tightening sequences show that the choice of sequence "
   "matters for the uniformity achieved %s. The underlying mechanism is elastic interaction: "
   "bolt loads change as neighbouring bolts are tightened. Experimentally determined "
   "interaction coefficients allowed a one-pass, tension-controlled bolt-up to reach "
   "\u00b12%% bolt-stress uniformity %s, and analytical models of the interaction have been "
   "used to optimise the loads applied in each pass %s. These studies optimise the sequence "
   "for joint quality; none treats the time taken by an automated tool to execute it, which "
   "depends on the path the tool must travel between bolts." % (
       C("pcc1"), C("brown2010"), C("kumakura2003"), C("bibel1992"), C("zhu2018", "zhu2019")))
H2("2.2 Torque control, preload scatter and torque\u2013angle monitoring")
Pp("Torque control infers preload from applied torque through the nut factor, and the "
   "torque\u2013tension relation is highly sensitive to thread and bearing friction %s. In "
   "gasketed pipe flanges tightened by torque control the scatter of axial bolt force is "
   "substantial and affects sealing performance %s, the torque coefficient and bolt preload "
   "are closely correlated %s, and finite-element studies find preload scatter due to elastic "
   "interaction difficult to eliminate under torque control %s. The accuracy of the assembly "
   "tool and the tightening strategy also affect clamp-force scatter %s. Monitoring the angle "
   "of turn alongside torque provides a second, friction-independent indication of clamp "
   "force, and torque\u2013angle formulations of the tightening process have been proposed "
   "and evaluated %s." % (C("nassar2007speed"), C("sawa2003"), C("omiya2016"), C("abid2013"),
                           C("persson2021"), C("nassar2007angle")))
H2("2.3 Robotic fastening and assembly data")
Pp("Robotic fastening research has addressed bolt insertion and tightening with tactile "
   "in-hand localisation %s, vision-based identification and pose estimation of nuts for a "
   "tightening robot %s, an adjustable spanner that engages bolts of different sizes "
   "automatically %s, and online vision inspection of screw-fastening quality %s. Digital "
   "records of assembly data and process traceability have been developed for complex "
   "products using digital-twin approaches %s. These works concern small fasteners, "
   "structured factory settings or single tightening operations; large multi-pass flange "
   "closures in hazardous areas have received little attention." % (
       C("nozu2018"), C("yibang2023"), C("ali2018"), C("martinez2020"), C("zhuang2021")))
H2("2.4 Heavy-component handling")
Pp("Gravity compensation of manipulators may use counterweights, springs or an auxiliary "
   "active force %s. For pneumatic actuation the available force is bounded by supply "
   "pressure and by standard cylinder sizes, which for the common ISO series reach a 320 mm "
   "bore %s; whether a multi-tonne head can be handled pneumatically is therefore a question "
   "of how much of the load a passive counterbalance carries." % (C("arakelian2016"),
                                                                  C("iso15552")))
H2("2.5 Automation in hazardous process environments")
Pp("Reviews of robotics in onshore and offshore oil and gas operations describe systems in "
   "which robots act while skilled operators retain cognitive decisions %s. Equipment in "
   "explosive atmospheres must follow the relevant protection concepts and certification "
   "requirements %s, and interlocks that protect against process hazards fall within "
   "functional-safety practice for the process industry %s. Model checking has been applied "
   "to verify PLC software %s, failure mode and effects analysis (FMEA) is the standard means "
   "of ranking design risks, with well-known limitations of the risk priority number %s, and "
   "discrete-event simulation is widely used to evaluate manufacturing system designs before "
   "they are built %s." % (C("shukla2016a", "shukla2016b"), C("singh2023", "iec60079", "atex"),
                           C("iec61511"), C("ovatman2016"), C("liu2013"), C("negahban2014")))
H2("2.6 Research gap")
Pp("The literature optimises flange tightening for joint quality, develops robotic tightening "
   "for small fasteners, and verifies PLC logic in general. What has not been done is to "
   "evaluate an integrated automated architecture for large multi-pass flange closures "
   "quantitatively: how long it takes as a function of its configuration, what preload "
   "accuracy its control strategy can actually deliver, whether its interlock logic is both "
   "safe and free of deadlock, and whether its head-handling concept is physically "
   "feasible. This paper addresses those four questions for one industrial case.")

# ==================================================================== 3
H("3 Case study, requirements and architecture")
H2("3.1 Industrial case")
Pp("The case is a four-drum delayed-coking unit. Each drum is approximately 28 m tall and "
   "6 m in diameter. Three bolted connections are opened and closed on every cycle: the top "
   "head (32 bolts, M48\u2013M64), the bottom head (48 bolts, M48\u2013M64) and the feed inlet "
   "(12 bolts, M36\u2013M48). Opening comprises unbolting, bolt storage, head removal and "
   "gasket removal; closing comprises face cleaning, gasket installation, head placement, "
   "bolt installation, multi-pass tightening, verification and a steam leak test. Hydraulic "
   "decoking, which occurs between opening and closing, is outside the scope. Flange drawings "
   "and head masses were not available; where the analysis needs them it treats them as "
   "swept parameters rather than assumed point values.", indent=False)
H2("3.2 Requirements")
TBL(["Requirement", "Target", "Source"],
    [["Bolts per operation", "92 (32 + 48 + 12)", "Case data"],
     ["Bolt sizes", "M36\u2013M64", "Case data"],
     ["Automated operations", "Breakout, removal, storage, re-presentation, installation, "
      "multi-pass tightening", "Design intent"],
     ["Tightening procedure", "Star passes at increasing torque + circular check passes", "[%s]" % C("pcc1")[1:-1]],
     ["Record per bolt", "Final torque and angle, pass/fail, time stamp", "Design intent"],
     ["Bolting time (92 bolts)", "1.5\u20132.5 h (reported manual: 4\u20136 h)", "Source proposal"],
     ["Complete operation (open + close, excluding decoking)", "Not stated in the source; "
      "evaluated here against the same 2.5 h", "This paper"],
     ["Personnel in hazard zone", "None during routine operation", "Design intent"],
     ["Tool motion", "Only with gas clearance confirmed", "Design intent"],
     ["Head motion", "Only with no bolt engaged", "Design intent"],
     ["CLOSED status", "Only after all bolts verified, leak test passed and operator sign-off",
      "Design intent"],
     ["Emergency stop", "From any state, local and remote", "Design intent"]],
    widths=[1.7, 3.1, 1.5])
CAP("Table 1. Functional requirements. 'Source proposal' marks the targets whose "
    "achievability this paper tests.")
H2("3.3 Architecture")
Pp("The architecture (Fig. 1) has four subsystems. (A) An orbital bolt runner is a circular "
   "rail fixed around each flange carrying a carriage with one or more pneumatic torque "
   "multipliers, each with a socket, reaction arm, and torque and angle transducers; with "
   "several heads they are spaced equally and act simultaneously. (B) A bolt-management "
   "carousel captures each bolt as it is removed, stores it in an indexed slot and presents "
   "it for reinstallation in the programmed order. (C) A counterbalanced articulated "
   "pneumatic manipulator engages, lifts, swings clear and locks each head. (D) A sequence "
   "and interlock controller enforces the tightening protocol, records torque and angle for "
   "every bolt, and exchanges start, stop and status signals with the plant control system. "
   "Pneumatic actuation is chosen because compressed air is available as plant "
   "infrastructure and avoids electrical ignition sources at the tool; the hazardous-area "
   "certification of the integrated system is discussed in Section 6.")
FIG("figure1_architecture")
CAP("Fig. 1 Architecture of the automated bolting and head-handling system for a "
    "three-flange drum. Subsystems A\u2013C are replicated at each flange (C at the two heads "
    "only); D is common.")

# ==================================================================== 4
H("4 Methods")
H2("4.1 Geometry")
Pp("Bolt i of a flange with N bolts sits at angle")
EQ("\u03b8\u1d62 = 2\u03c0i / N,   i = 0, 1, \u2026, N \u2212 1")
Pp("on a bolt circle of diameter D. Because D was not available, it is bracketed by a "
   "circumferential spacing rule s = k\u00b7d, with d the nominal bolt diameter and k between "
   "%.1f and %.1f (wrench clearance sets the lower end, gasket-load uniformity the upper):"
   % tuple(R["k_spacing"]), indent=False)
eq_D = EQ("D = N k d / \u03c0")
Pp("giving D from %.2f to %.2f m for the bottom head, %.2f to %.2f m for the top head and "
   "%.2f to %.2f m for the feed inlet. The spacing factor k is sampled in the Monte Carlo "
   "analysis, so its effect appears in the sensitivity results rather than being fixed."
   % tuple(R["sequences"]["bottom head"]["D_range_m"] + R["sequences"]["top head"]["D_range_m"]
           + R["sequences"]["feed inlet"]["D_range_m"]), indent=False)
H2("4.2 Tightening sequence and rail travel")
Pp("The legacy star pattern for N divisible by four is generated as groups of four bolts at "
   "90\u00b0 spacing, with the groups visited in bit-reversed order, which reproduces the "
   "cross-pattern family of the legacy sequence %s. A circular pass visits bolts in angular "
   "order. For a pass visiting stations j\u2081, j\u2082, \u2026 the carriage travels the "
   "shortest way round the rail between consecutive stations," % C("pcc1"))
EQ("L = (\u03c0D / N) \u03a3\u2096 \u03b4(j\u2096, j\u2096\u208a\u2081),   "
   "\u03b4(a, b) = min((b \u2212 a) mod M, M \u2212 (b \u2212 a) mod M)")
Pp("where M = N/r. With r equally spaced heads acting together, bolts j, j + N/r, \u2026 are "
   "tightened at one station, so a pass visits M stations, and the carriage reaches the next "
   "station by moving whichever head is nearest, i.e. the shortest way on a ring of M "
   "positions (for r = 1, M = N). Fig. 2 shows the 48-bolt star sequence and the resulting "
   "travel per pass: a single head covers %.0f bolt-circle diameters per star pass on the "
   "bottom head, against %.1f for a circular pass and %.1f with four heads."
   % (R["sequences"]["bottom head"]["travel_star_perD"],
      R["sequences"]["bottom head"]["travel_circ_perD"],
      R["sequences_travel_4heads_perD"]["bottom head"]), indent=False)
Pp("Much of that travel might be regarded as wasted motion, so the tool path is also "
   "optimised. Let the quartets g = 0, \u2026, q \u2212 1 (q = N/4) be visited in an order "
   "\u03c0; each quartet is worked as g, g + 2q, g + q, g + 3q, so that opposite bolts are "
   "loaded consecutively as the legacy procedure requires. The optimised order solves")
EQ("min over \u03c0 of \u03a3\u2096 c(\u03c0\u2096, \u03c0\u2096\u208a\u2081)   subject to   "
   "\u03b4_q(\u03c0\u2096, \u03c0\u2096\u208a\u2081) \u2265 s_min for all k")
Pp("where c is the carriage travel from the last bolt of one quartet to the first of the "
   "next plus the travel within the next quartet, and \u03b4_q is the separation of two "
   "quartets on the ring of q positions. The load-spreading constraint requires every "
   "consecutive pair of quartets to be at least as far apart as the closest consecutive "
   "pair in the legacy sequence (s_min = %d for the top and bottom heads), so the optimised "
   "path never concentrates load more locally than the legacy path does at any step. The "
   "problem is a constrained shortest Hamiltonian path on at most q = 12 nodes and is solved "
   "exactly by dynamic programming (Held\u2013Karp). Whether a sequence that satisfies this "
   "geometric constraint also preserves preload uniformity under elastic interaction must be "
   "established experimentally %s." % (SQ["bottom head"]["legacy_min_separation"],
                                       C("zhu2018", "zhu2019")), indent=False)
FIG("figure2_sequence")
CAP("Fig. 2 (a) Legacy star sequence for the 48-bolt bottom head; numbers give the order, "
    "red arrows the first eight carriage moves, orange one group of four bolts tightened "
    "together by a four-head carriage. (b) Rail travel per pass, in bolt-circle diameters, "
    "for each flange and configuration.")
H2("4.3 Cycle-time model")
Pp("One pass over a flange takes")
EQ("t_pass = (N/r)(t_align + t_engage + t_tool + t_release + t_verify + t_retry + t_review) "
   "+ L/v + m\u00b7t_move")
Pp("where t_tool is the torque, breakout or run-on/run-off time of the pass, v the carriage "
   "speed, m the number of indexing moves and t_move a per-move acceleration and settling "
   "overhead. A station is complete only when all r heads are complete, so the expected "
   "retry and operator-review delays per station are", indent=False)
EQ("t_retry = [1 \u2212 (1 \u2212 p_retry)\u02b3](t_engage + t_tool),   "
   "t_review = [1 \u2212 (1 \u2212 p_op)\u02b3] t_op")
Pp("Run-on and run-off take the time to turn the nut through its engagement length, "
   "t_spin = 60\u00b7(1.5d/P)/n_spin, with P the ISO coarse pitch %s and n_spin the spin speed "
   "in rev/min. Opening is one staged breakout pass in star order, a circular run-off pass and "
   "transfer of every bolt to the carousel; closing is transfer from the carousel, a circular "
   "run-on pass, three star passes at increasing torque and one to three circular check "
   "passes, following the legacy procedure %s. Bolt transfer between the flange and the "
   "carousel takes N\u00b7t_transfer/c, where c is the number of transfer channels working in "
   "parallel: c = 1 for a single carousel serving one bolt at a time, c = r when each torque "
   "head has its own feed point. Flanges are either worked in parallel, each by its own "
   "runner, or in series." % (C("iso261"), C("pcc1")), indent=False)
Pp("Two cycle times are distinguished. The bolting time is the time to unbolt and re-bolt "
   "all 92 bolts, the quantity to which the source proposal's 1.5\u20132.5 h target referred:",
   indent=False)
EQ("T_bolting = max_f T_open,f + max_f T_close,f")
Pp("The complete operation adds atmosphere confirmation, head handling, gasket and face work "
   "and the steam leak test. In a phase-synchronised schedule every flange finishes a phase "
   "before the next begins; in a pipelined schedule each flange runs its own chain and only "
   "the leak test waits for all of them:", indent=False)
EQ("T_complete = t_atmos + max_f(T_open,f + t_head,f) + max_f(t_gasket + t_head,f + T_close,f) "
   "+ t_leak")
Pp("with t_head,f = 0 for the feed inlet. Hydraulic decoking, which separates opening from "
   "closing, is excluded from both.", indent=False)
STEP_LBL = [("v_carriage", "Carriage speed", "m/s"), ("t_move", "Indexing move overhead", "s"),
            ("t_align", "Alignment to nut", "s"), ("t_engage", "Socket engagement", "s"),
            ("t_torque", "Torque pass, one bolt", "s"), ("t_breakout", "Breakout, one bolt", "s"),
            ("t_release", "Reaction release and retract", "s"), ("t_verify", "Torque\u2013angle check", "s"),
            ("rpm_spin", "Run-on/run-off speed", "rev/min"), ("t_transfer", "Carousel transfer, one bolt", "s"),
            ("p_retry", "P(automatic retry) per bolt-pass", "\u2013"),
            ("p_operator", "P(operator review) per bolt-pass", "\u2013"),
            ("t_operator", "Operator review duration", "s"), ("check_passes", "Circular check passes", "\u2013"),
            ("t_head", "Head handling (each way)", "s"), ("t_gasket", "Gasket and face work, per flange", "s"),
            ("t_atmos", "Atmosphere confirmation", "s"), ("t_leak", "Steam leak test", "s")]
TBL(["Parameter", "Range (uniform)", "Unit"],
    [[lab, "%g \u2013 %g" % tuple(R["step_ranges_s"][k]), u] for k, lab, u in STEP_LBL],
    widths=[3.0, 1.8, 1.0])
CAP("Table 2. Assumed step-time ranges. None is a measurement; the ranges are the inputs "
    "whose influence the sensitivity analysis ranks, and the validation plan (Section 7) "
    "specifies how each is to be measured.")
H2("4.4 Monte Carlo simulation and sensitivity")
Pp("All inputs in Table 2, the spacing factor k and the position of each bolt size within "
   "its range are sampled independently %d times with a fixed seed (%d). Six configurations "
   "are compared: one head with flanges in parallel, in series, and with circular passes "
   "only; two heads; four heads; and four heads with two instead of three star passes, as an "
   "example of a reduced-pass procedure of the kind permitted as an alternative to the "
   "legacy sequence %s. Sensitivity is measured by the Spearman rank correlation between each "
   "sampled input and the total bolting time. A best case, with every step at the fast end of "
   "its range, bounds what each configuration could achieve. A second study then removes the "
   "bottlenecks of the four-head design one at a time, on the same samples: (A) four heads "
   "and a single carousel; (B) one transfer channel per head; (C) two star passes; (D) the "
   "optimised tool path; (E) pipelined servicing. Three uncertainty models are compared to "
   "test whether the conclusions depend on the distributional assumptions (Section 5.6). "
   "Spearman rank correlation is used as a computationally transparent screening measure of "
   "sensitivity; variance-based (Sobol) analysis is deferred until empirical parameter "
   "distributions are available from the validation programme, since apportioning variance "
   "among assumed distributions would add precision without adding information. "
   "Discrete-event simulation of this kind is "
   "standard for evaluating manufacturing designs before they are built %s."
   % (R["n_mc"], R["seed"], C("brown2010"), C("negahban2014")))
H2("4.5 Preload model")
Pp("The torque needed for a target preload F in a bolt of diameter d is")
EQ("T = K F d")
Pp("with K the nut factor. The target is taken as half the specified minimum yield strength "
   "of SA-193 B7 bolting, 724 MPa %s, acting on the tensile stress area %s"
   % (C("astm193"), C("iso898")), indent=False)
EQ("A_s = (\u03c0/4)(d \u2212 0.9382P)\u00b2")
Pp("The achieved preload under torque control is F = T(1 + \u03b5_tool)/(K d), with tool "
   "error \u03b5_tool uniform within the specified accuracy and K normally distributed with "
   "coefficient of variation CV_K. To first order", indent=False)
EQ("CV_F \u2248 (CV_tool\u00b2 + CV_K\u00b2)^\u00bd")
Pp("so the benefit of a tighter torque specification depends on its size relative to "
   "CV_K. Torque\u2013angle monitoring is modelled as an independent estimate of preload from "
   "the angle turned after snug, with error CV_\u03b8 representing joint and gasket stiffness "
   "uncertainty; bolts whose estimate lies outside \u00b110% of target are re-torqued to the "
   "angle-corrected value. Each case uses 200,000 samples.", indent=False)
H2("4.6 Head manipulator force budget")
Pp("For a head of mass M_h lifted with safety factor SF = 1.5, the pneumatic force needed "
   "when a passive counterbalance carries a fraction c of the weight is")
EQ("F_p = SF(1 \u2212 c) M_h g")
Pp("and the bore of a single cylinder at supply pressure p = 0.6 MPa and force efficiency "
   "\u03b7 = 0.9 is b = [4F_p/(\u03c0p\u03b7)]^\u00bd. The minimum counterbalance fraction "
   "for the largest standard bore (320 mm) %s follows by setting b = 320 mm. Because the head "
   "mass is not known, M_h is swept from 1 to 20 t." % C("iso15552"), indent=False)
H2("4.7 Control logic and its verification")
Pp("The controller is written as an explicit transition function over 13 modes (Fig. 3) and "
   "a state comprising the number of engaged bolts, the number verified, the leak-test result "
   "and the phase to resume after an interruption. Nine Boolean inputs (gas clearance, "
   "emergency stop, torque pass, angle pass, head aligned, leak pass, start, confirm, reset) "
   "are applied in every combination from every reachable state, by breadth-first "
   "enumeration, and seven safety invariants are checked on every transition: I1 tool motion "
   "begins only with gas clearance; I2 the head moves only with no bolt engaged; I3 the "
   "sequence advances by at most one bolt and only on a torque and angle pass; I4 CLOSED "
   "follows only a passed leak test with all bolts verified; I5 CLOSED requires operator "
   "sign-off; I6 an emergency stop forces the safe state from any state; I7 loss of gas "
   "clearance halts tool motion. A liveness check then identifies reachable states from "
   "which no input sequence can reach a completed state (OPEN or CLOSED). To test that the "
   "invariants have discriminating power, each interlock is removed in turn and the check "
   "repeated; a check that no such mutant fails would be testing nothing. This is "
   "explicit-state model checking of the logic as specified %s; it does not verify a PLC "
   "implementation." % C("ovatman2016"))
FIG("figure3_state_machine")
CAP("Fig. 3 Sequence and interlock state machine. Dashed transitions are interlock "
    "responses; the resume rule on reset (bottom right) was added after the liveness check "
    "found deadlocks in the logic as first specified.")
H2("4.8 FMEA, automation coverage and exposure")
Pp("Twelve failure modes are scored for severity S, occurrence O (with the stated control) "
   "and detection D on 1\u201310 scales, and ranked by RPN = S\u00b7O\u00b7D %s. The scores "
   "are the author's engineering judgement, not field data. Automation coverage is"
   % C("liu2013"))
EQ("A = \u03a3 w\u1d62a\u1d62 / \u03a3 w\u1d62")
Pp("with a\u1d62 = 1 for fully automatic, 0.5 for operator-initiated and 0 for manual "
   "operations, weighted by modelled task duration w\u1d62. Hazard-zone exposure is counted "
   "in person-hours per operation.", indent=False)
Pp("Implementation note: the models were written in Python with the assistance of a "
   "generative AI tool (see Declarations); the author specified, reviewed and checked every "
   "model and result.", indent=False, size=10)

# ==================================================================== 5
H("5 Results")
H2("5.1 Torque requirement")
TBL(["Bolt", "A\u209b (mm\u00b2)", "Target preload (kN)", "Torque, K = 0.16 (kN\u00b7m)",
     "Torque, K = 0.20 (kN\u00b7m)"],
    [["M%s" % d, "%.0f" % v["As_mm2"], "%.0f" % v["F_kN"], "%.1f" % v["T_K0.16"], "%.1f" % v["T_K0.20"]]
     for d, v in tq.items()])
CAP("Table 3. Target preload at half the specified minimum yield strength of SA-193 B7 and "
    "the corresponding torque for two nut factors.")
Pp("Table 3 gives the torques the runner must deliver: %.1f\u2013%.1f kN\u00b7m for the M64 "
   "bolts of the heads and %.1f\u2013%.1f kN\u00b7m for M36 on the feed inlet. These are "
   "within the range of torque multipliers, and they set the reaction load the carriage and "
   "rail must resist at every station." % (tq["64"]["T_K0.16"], tq["64"]["T_K0.20"],
                                          tq["36"]["T_K0.16"], tq["36"]["T_K0.20"]))
H2("5.2 Cycle time")
TBL(["Configuration", "Opening (h)", "Closing (h)", "Bolting total (h)", "P(\u2264 2.5 h)",
     "Best case (h)"],
    [[k, "%.2f" % v["bolting_open"]["median"], "%.2f" % v["bolting_close"]["median"],
      "%.2f [%.2f\u2013%.2f]" % (v["bolting_total"]["median"], v["bolting_total"]["p05"],
                                  v["bolting_total"]["p95"]),
      "%.2f" % v["p_bolting_le_2.5h"], "%.2f" % best[k]] for k, v in cyc.items()],
    widths=[2.0, 0.8, 0.8, 1.4, 0.8, 0.8])
CAP("Table 4. Modelled bolting time for all 92 bolts (median, with 5th\u201395th percentiles), "
    "probability of meeting the 2.5 h target, and the best case with every step at its fast "
    "limit. Opening and closing exclude head handling, gasket work and leak test.")
FIG("figure4_cycle_time")
CAP("Fig. 4 Distribution of total bolting time by configuration, with the 1.5\u20132.5 h "
    "target of the source proposal and the reported 4\u20136 h manual duration.")
Pp("With one head per flange and the flanges worked in parallel, bolting takes %s, set by "
   "the bottom head. Working the flanges in series nearly doubles this to %.1f h, and using "
   "only circular passes, which would not satisfy the legacy procedure, still leaves %.1f h. "
   "Most importantly, within the investigated parameter bounds the single-head configuration "
   "cannot meet the 2.5 h target even in its "
   "best case of %.2f h. A second head per flange halves the time to %.1f h. Only four "
   "synchronised heads bring the median within the target, at %s, and even then the target "
   "is met in %.0f%% of samples; four heads with one fewer star pass raise this to %.0f%%. "
   "Including head handling, gasket work and the leak test, the full opening and closing "
   "operation takes %.1f h with one head and %.1f h with four."
   % (mr(b1["bolting_total"]), bs["bolting_total"]["median"], bc["bolting_total"]["median"],
      best["1 head, parallel flanges"], b2["bolting_total"]["median"], mr(b4["bolting_total"]),
      100 * b4["p_bolting_le_2.5h"], 100 * b4r["p_bolting_le_2.5h"],
      b1["full_total"]["median"], b4["full_total"]["median"]))
H2("5.3 Admissibility of the cycle-time target")
Pp("The legacy procedure applies each bolt roughly eight times per open\u2013close operation "
   "(one breakout, one run-off, one run-on, three star passes and a median of two check "
   "passes), giving %d bolt-passes on the bottom head and %d in total. Meeting 2.5 h with "
   "flanges in parallel therefore allows %.0f s per bolt-pass on the bottom head, and "
   "%.0f s if the flanges share one runner in series. The modelled median is %.0f s per "
   "bolt-pass. The target is thus not a matter of faster tools: alignment, engagement, "
   "torque, release and verification of a large bolt cannot plausibly fit in about 20 s. It "
   "is achievable only by working several bolts at once, which is what a four-head carriage "
   "does, and which is also what a manual crew of four to six does."
   % (adm["bolt_passes_per_flange"]["bottom head"], adm["bolt_passes_total"],
      adm["budget_s_per_bolt_pass_parallel_2.5h"], adm["budget_s_per_bolt_pass_serial_2.5h"],
      adm["modelled_median_s_per_bolt_pass_bottom"]))
H2("5.4 Sensitivity")
top1 = list(sens1.items())[:3]
top4 = list(sens4.items())[:3]
Pp("With one head (Fig. 5a) the time is driven by carriage speed (\u03c1 = %.2f), torque time "
   "per pass (%.2f) and the number of check passes (%.2f): the single head spends much of the "
   "pass travelling the star pattern. With four heads (Fig. 5b) travel almost disappears and "
   "the ranking changes: carousel transfer becomes the strongest driver (\u03c1 = %.2f), "
   "followed by check passes (%.2f) and torque time (%.2f), and the probability of an "
   "operator review rises in importance because any one of four heads can stall a station. "
   "Carousel transfer accounts for %.0f%% of bottom-head bolting time in the four-head "
   "configuration. A design with four heads must therefore also increase carousel throughput, "
   "for example with a feed point per head; otherwise the carousel, not the torque tool, sets "
   "the pace."
   % (sens1["v_carriage"], sens1["t_torque"], sens1["check_passes"], sens4["t_transfer"],
      sens4["check_passes"], sens4["t_torque"], 100 * R["carousel_share_4heads"]))
FIG("figure5_sensitivity")
CAP("Fig. 5 Spearman rank correlation of each sampled input with total bolting time: (a) one "
    "head per flange, (b) four heads per flange. The nine strongest inputs are shown.")
H2("5.5 Throughput-optimised architecture")
ROWS5 = [("One head (baseline)", b1), ("A  Four heads, single carousel", OA),
         ("B  + four transfer channels", OB), ("C  + two star passes", OC),
         ("D  + optimised tool path", OD), ("E  + pipelined servicing", OE)]


def rng_(x):
    return "%.2f [%.2f–%.2f]" % (x["median"], x["p05"], x["p95"])


TBL(["Configuration", "Bolting (h)", "P(bolting ≤ 2.5 h)", "Complete operation (h)",
     "P(complete ≤ 2.5 h)", "P(complete ≤ 4 h)"],
    [[nm, rng_(v["bolting_total"]),
      "%.2f" % v.get("p_bolting_le_2.5h", b1["p_bolting_le_2.5h"]), rng_(v["full_total"]),
      "%.2f" % v.get("p_full_le_2.5h", 0.0),
      ("%.2f" % v["p_full_le_4.0h"]) if "p_full_le_4.0h" in v else "0.00"]
     for nm, v in ROWS5], widths=[2.0, 1.2, 0.8, 1.3, 0.8, 0.8])
CAP("Table 5. Removing the bottlenecks of the automated architecture one at a time (median "
    "and 5th–95th percentiles, same Monte Carlo samples throughout). Each row adds one "
    "change to the row above.")
Pp("Table 5 and Fig. 6 remove the bottlenecks in turn. Giving each of the four heads its own "
   "bolt-transfer channel, for example a carousel with four access stations, cuts bolting from "
   "%.2f to %.2f h and raises the probability of meeting 2.5 h from %.2f to %.2f. This is a "
   "change to the machine alone, with the legacy tightening procedure unaltered, and it is the "
   "configuration the analysis recommends: four heads and four transfer channels meet the "
   "bolting target without any change to how the joint is assembled. Removing one star pass "
   "brings bolting further to %.2f h, inside the target in every sample, but this is a change "
   "to the joint-assembly procedure rather than to the machine. It is an optional process "
   "improvement, admissible only if a reduced-pass procedure of the kind described as an "
   "alternative to the legacy sequence %s is shown to give equivalent preload uniformity on "
   "the actual joint (Section 7); the architecture does not depend on it."
   % (OA["bolting_total"]["median"], OB["bolting_total"]["median"], OA["p_bolting_le_2.5h"],
      OB["p_bolting_le_2.5h"], OC["bolting_total"]["median"], C("brown2010")))
Pp("The exactly optimised tool path adds little, and it is reported for what it shows rather "
   "than as part of the recommended design. With four heads it cuts carriage travel per "
   "star pass on the bottom head from %.2f to %.2f bolt-circle diameters, but travel is by "
   "then a small part of the cycle and bolting falls only from %.2f to %.2f h. With one head, "
   "where travel matters, the optimum reduces travel per star pass only from %.1f to %.1f "
   "diameters and bolting from %.2f to %.2f h. The reason is structural: most of the "
   "single-head travel is incurred within each quartet, because opposite bolts, half a turn "
   "apart, must be loaded consecutively, and no ordering of the quartets can remove that. The "
   "wasted motion of Fig. 2 is thus a consequence of serving a symmetric loading rule with one "
   "tool, and the remedy is more tools rather than a better path. Pipelining the servicing "
   "steps changes nothing (%.2f h in both D and E) because the bottom head, with the most "
   "bolts and a head to handle, is on the critical path in every phase, so the other "
   "flanges' chains finish within it."
   % (SQ["bottom head"]["legacy_r4"], SQ["bottom head"]["optimised_r4"],
      OC["bolting_total"]["median"], OD["bolting_total"]["median"],
      SQ["bottom head"]["legacy_r1"], SQ["bottom head"]["optimised_r1"],
      b1["bolting_total"]["median"], O1["bolting_total"]["median"],
      OE["full_total"]["median"]))
Pp("The complete operation is a different matter. Even in configuration E it takes %.2f h "
   "(90%% interval %.2f–%.2f h) and meets 2.5 h in none of the %d samples. Of the median, "
   "%.2f h is bolting and %.2f h is servicing: head handling in both directions (%.2f h), the "
   "leak test (%.2f h), gasket and face work (%.2f h) and atmosphere confirmation (%.2f h). "
   "The 2.5 h target is therefore achievable for bolting, which is what the source proposal "
   "stated, but not for the complete operation; once bolting is parallelised, head handling "
   "and the leak test become the next bottlenecks."
   % (OE["full_total"]["median"], OE["full_total"]["p05"], OE["full_total"]["p95"], R["n_mc"],
      COMP["bolting"], COMP["non_bolting_median_h"], COMP["head_both_ways"], COMP["leak_test"],
      COMP["gasket"], COMP["atmos"]))
SE = R["servicing_envelope"]
Pp("That observation can be turned into a performance requirement for the servicing "
   "subsystems. For a complete operation of at most 2.5 h, the servicing steps must satisfy")
EQ_SERVICE = EQ("t_atmos + 2t_head + t_gasket + t_leak ≤ 2.5 h − T_bolting")
Pp("With the recommended configuration B (median bolting %.2f h) the servicing budget is "
   "%.2f h, or %.0f min; with the optional reduced-pass configuration C (%.2f h) it is %.2f h, "
   "or %.0f min. The assumed servicing ranges give %.2f h at the median and %.2f h at best. "
   "With configuration C, a 2.5 h complete operation would need, for example, head handling "
   "of about 10 min each way, gasket and face work of about 10 min, a leak test of about "
   "15 min and atmosphere confirmation under 5 min, about 50 min in all; with B the same "
   "steps would have to fit in about 40 min. Either is roughly a halving of every servicing "
   "step. "
   "Equation (%d) is the envelope any faster head-handling or leak-test design would have to "
   "meet; no such design is evaluated here."
   % (SE["B + 4 transfer channels"]["bolting_median_h"], SE["B + 4 transfer channels"]["servicing_budget_h"],
      SE["B + 4 transfer channels"]["servicing_budget_min"], SE["C + 2 star passes"]["bolting_median_h"],
      SE["C + 2 star passes"]["servicing_budget_h"], SE["C + 2 star passes"]["servicing_budget_min"],
      SE["assumed_median_h"], SE["assumed_min_h"], EQN[0]), indent=False)
FIG("figure6_optimisation")
CAP("Fig. 6 (a) Bolting time and (b) complete-operation time as the bottlenecks are removed "
    "one at a time (Table 5). Shading marks 2.5 h; labels give the median and the probability "
    "of meeting 2.5 h.")
H2("5.6 Robustness to the uncertainty model")
RBU = R["robustness"]
dnames = list(RBU)
cn = ["1 head", "2 heads", "A 4 heads", "B + 4 channels", "C + 2 star passes"]
TBL(["Configuration"] + ["%s: median (h) / P(≤ 2.5 h)" % d for d in dnames],
    [[c_] + ["%.2f / %.2f" % (RBU[d][c_]["median"], RBU[d][c_]["p_le_2.5h"]) for d in dnames]
     for c_ in cn] + [["Ranking 1 > 2 > A > B holds in"] +
                      ["%.1f%% of samples" % (100 * RBU[d]["ranking_holds_fraction"]) for d in dnames]],
    widths=[1.6, 1.6, 1.6, 1.6])
CAP("Table 6. Bolting time under three uncertainty models: the baseline independent uniform "
    "distributions, independent symmetric triangular distributions with the mode at the "
    "midpoint of each range, and uniform margins with a correlation of 0.6, through a Gaussian "
    "copula, among alignment, engagement and torque times, retry and review probabilities and "
    "bolt size.")
Pp("The step-time ranges of Table 2 are assumptions, and so are the uniform distributions and "
   "the independence with which they are sampled. It is plausible that a large, badly "
   "aligned bolt takes longer to align, engage and torque and is also likelier to need a "
   "retry or review, so that those inputs move together. Table 6 repeats the "
   "analysis under two alternatives: triangular distributions centred on the nominal value, "
   "and a pessimistic case in which those six inputs are positively correlated. The medians "
   "barely move, because the ranges are unchanged; the tails do. In the correlated case the "
   "probability that the recommended configuration B meets 2.5 h falls from %.2f to %.2f and "
   "its 95th percentile rises from %.2f to %.2f h, whereas under triangular distributions it "
   "rises to %.3f. What does not change is the architectural ranking: one head is slower "
   "than two, two slower than four, and four heads with a single carousel slower than four "
   "with four transfer channels, in every one of the %d samples under all three models. The "
   "paper's central claim is that ranking, not the particular value of %.2f h."
   % (RBU[dnames[0]]["B + 4 channels"]["p_le_2.5h"], RBU[dnames[2]]["B + 4 channels"]["p_le_2.5h"],
      RBU[dnames[0]]["B + 4 channels"]["p95"], RBU[dnames[2]]["B + 4 channels"]["p95"],
      RBU[dnames[1]]["B + 4 channels"]["p_le_2.5h"], R["n_mc"],
      RBU[dnames[0]]["B + 4 channels"]["median"]))
FM = R["feasibility_map"]
Pp("Fig. 7 generalises the discrete configurations into a design feasibility map over the "
   "number of torque heads and transfer channels per flange, for the legacy three-star-pass "
   "procedure and a two-pass alternative. Heads are limited to 1, 2 and 4 because they must "
   "be equally spaced on all three flanges, and 32, 48 and 12 bolts are all divisible only "
   "by those. The feasible region, where bolting meets 2.5 h in at least 90%% of samples, "
   "requires four heads under either procedure: with four heads, two transfer channels give "
   "%.2f under the legacy procedure, and the two-pass procedure reaches %.2f, just short of "
   "the threshold, even with a single channel. With two heads no combination exceeds %.2f. Adding channels without adding heads "
   "achieves almost nothing, because transfer is a bottleneck only once torque application "
   "has been parallelised."
   % (FM["3,4,2"]["p_le_2.5h"], FM["2,4,1"]["p_le_2.5h"],
      max(v["p_le_2.5h"] for k, v in FM.items() if k.split(",")[1] == "2")))
FIG("figure7_feasibility", width=6.0)
CAP("Fig. 7 Design feasibility map: probability that bolting meets 2.5 h (colour and upper "
    "number) and median bolting time (lower number) against torque heads and bolt-transfer "
    "channels per flange, for (a) the legacy three-star-pass procedure and (b) a two-pass "
    "alternative that would require qualification.")
H2("5.7 Preload")
g = grid
Pp("Fig. 8a shows that the \u00b12%% torque accuracy stated for the system has almost no "
   "effect on preload scatter. With CV_K = 10%%, preload CV is %.1f%% for a \u00b12%% tool, "
   "%.1f%% for \u00b15%% and %.1f%% for \u00b110%%; %.0f%% of bolts fall outside \u00b110%% of "
   "target even with the \u00b12%% tool. Preload scatter is set by the nut factor, as the "
   "torque\u2013tension literature leads one to expect %s. The angle channel is what changes "
   "the outcome (Fig. 8b): with an angle-based preload estimate of 3%% uncertainty and "
   "re-torque of bolts outside \u00b110%%, the fraction outside falls from %.0f%% to %.0f%% "
   "(CV_K = 10%%) and from %.0f%% to %.0f%% (CV_K = 15%%), at the cost of re-torquing %.0f%% "
   "and %.0f%% of bolts respectively. With a 5%% angle-estimate uncertainty the residual is "
   "about %.0f%%; with 8%% the benefit largely disappears. The value of torque\u2013angle "
   "monitoring therefore depends on how well joint stiffness is known, which the validation "
   "plan measures."
   % (100 * g[(0.1, 0.02)]["preload_cv"], 100 * g[(0.1, 0.05)]["preload_cv"],
      100 * g[(0.1, 0.1)]["preload_cv"], 100 * g[(0.1, 0.02)]["frac_outside_10pct"],
      C("nassar2007speed", "sawa2003", "omiya2016"),
      100 * ta[(0.1, 0.03)]["outside10_before"], 100 * ta[(0.1, 0.03)]["outside10_after"],
      100 * ta[(0.15, 0.03)]["outside10_before"], 100 * ta[(0.15, 0.03)]["outside10_after"],
      100 * ta[(0.1, 0.03)]["flagged_frac"], 100 * ta[(0.15, 0.03)]["flagged_frac"],
      100 * ta[(0.1, 0.05)]["outside10_after"]))
FIG("figure8_preload")
CAP("Fig. 8 (a) Achieved preload coefficient of variation against nut-factor variation for "
    "three tool accuracies. (b) Fraction of bolts outside \u00b110% of target under torque "
    "control alone and with torque\u2013angle re-torque, for two nut-factor scatters and three "
    "angle-estimate uncertainties.")
H2("5.8 Head handling")
Pp("Without counterbalance, lifting a 5 t head at SF 1.5 from 6 bar air needs a %.0f mm "
   "cylinder and a 10 t head %.0f mm, beyond the largest standard bore. The minimum "
   "counterbalance fraction that keeps a single cylinder within 320 mm is %.0f%% for 5 t, "
   "%.0f%% for 10 t, %.0f%% for 15 t and %.0f%% for 20 t (Fig. 9). A pneumatic manipulator "
   "is therefore feasible only as a counterbalanced device in which air supplies the balance "
   "correction and motion, not the lift, consistent with passive gravity compensation "
   "%s. Compressibility also makes a pneumatic axis compliant, so the final seating of the "
   "head must rely on mechanical guides and alignment pins, and holding must rely on a "
   "spring-applied mechanical lock rather than on air pressure."
   % (man[5]["bore_mm_c0.00"], man[10]["bore_mm_c0.00"], 100 * man[5]["min_counterbalance"],
      100 * man[10]["min_counterbalance"], 100 * man[15]["min_counterbalance"],
      100 * man[20]["min_counterbalance"], C("arakelian2016")))
FIG("figure9_manipulator", width=5.2)
CAP("Fig. 9 Single-cylinder bore needed to lift a head against its unbalanced weight, as a "
    "function of head mass and counterbalance fraction. Head mass was not available and is "
    "swept.")
H2("5.9 Control logic")
Pp("With the full interlock set and the resume rule, exhaustive exploration finds no safety "
   "violation: %d reachable states and %d checked transitions for a 12-bolt flange (%d and "
   "%d for 4 bolts), with all 13 modes reachable and no deadlocked state. Every mutant is "
   "detected (Table 7), so each invariant is capable of failing. The liveness check, however, "
   "found a defect in the logic as first specified. Without an explicit rule for resuming an "
   "interrupted phase, %d of %d reachable states for a 12-bolt flange (%d of %d for 4 bolts) "
   "can never reach OPEN or CLOSED: an emergency stop part-way through unbolting or "
   "tightening leaves a partially bolted joint from which the prose interlocks provide no "
   "way forward, so the operator would have to leave the automated system. Adding the rule "
   "that atmosphere confirmation routes to the interrupted phase \u2014 continue unbolting, "
   "resume tightening at the next unverified bolt, repeat the leak test or await sign-off "
   "\u2014 removes every deadlock while preserving all seven invariants."
   % (fd12["states"], fd12["transitions"], fd4["states"], fd4["transitions"],
      as12["deadlocked_states"], as12["states"], as4["deadlocked_states"], as4["states"]))
MUT = {"esd": "Emergency stop", "gas_before_unbolt": "Gas clearance before tool motion",
       "gas_hold": "Hold on loss of gas clearance", "leak_before_closed": "Leak test before CLOSED",
       "no_head_move_engaged": "No head motion with bolts engaged",
       "sequence_lock": "Sequence advances only on torque AND angle pass",
       "signoff_before_closed": "Operator sign-off before CLOSED"}
TBL(["Interlock removed", "Invariant violated", "Shortest counterexample (transitions)"],
    [[MUT[k], v[0][0], str(v[0][1])] for k, v in V["mutants"].items()], widths=[2.3, 2.8, 1.2])
CAP("Table 7. Mutation test of the interlock logic (4-bolt flange). Each interlock removed "
    "in isolation is caught by the invariant it protects.")
H2("5.10 Failure modes")
Pp("The highest-ranked failure mode (Fig. 10, Table 8) is an out-of-window preload that goes "
   "undetected (RPN %d), which follows directly from Section 5.7: torque control alone cannot "
   "see a nut-factor outlier, and detection depends on the angle channel. Second is failure "
   "of gas detection to register residual hydrocarbons (RPN %d, severity 10), which is why "
   "the design specifies voted redundant detectors and a pre-cycle functional test. The "
   "failure modes with the highest severity but lower RPN \u2014 head release with a bolt "
   "engaged and loss of air under load \u2014 are those the interlocks and the mechanical lock "
   "address by design; their low RPN reflects the control, not the absence of the hazard."
   % (fm[0]["RPN"], fm[1]["RPN"]))
FIG("figure10_fmea")
CAP("Fig. 10 Design FMEA ranked by risk priority number; colour indicates severity. Scores "
    "are the author's judgement and are listed with their controls in Table 8.")
TBL(["Failure mode", "Effect", "S", "O", "D", "RPN", "Control"],
    [[f["mode"], f["effect"], f["S"], f["O"], f["D"], f["RPN"], f["control"]] for f in fm],
    fs=7.5, widths=[1.6, 1.3, 0.3, 0.3, 0.3, 0.4, 2.0])
CAP("Table 8. Design FMEA (1\u201310 scales; D = 10 is undetectable).")
H2("5.11 Automation coverage and exposure")
TBL(["Operation", "Modelled duration (h)", "Automation level a\u1d62"],
    [[t["task"], "%.2f" % t["hours"], "%.1f" % t["a"]] for t in ex["tasks"]], widths=[3.2, 1.5, 1.5])
CAP("Table 9. Operations, modelled durations (one-head configuration, medians) and "
    "automation levels used in the coverage metric.")
Pp("Weighted by modelled duration, automation coverage is A = %.2f (Table 9). The reported "
   "manual procedure places four to six workers in the hazard zone for four to six hours, "
   "%d\u2013%d person-hours per operation; with the automated system the routine value is "
   "zero, rising to at most %.2f person-hours if gasket cassette loading has to be done at the "
   "flange rather than from outside the zone. The exposure reduction is the least uncertain "
   "benefit of the architecture: it does not depend on the step times that dominate the "
   "cycle-time uncertainty."
   % (ex["coverage"], ex["manual_person_hours"][0], ex["manual_person_hours"][1],
      ex["pabhs_person_hours_in_zone"][1]))

# ==================================================================== 6
H("6 Discussion")
H2("6.1 What the targets require")
Pp("The analysis supports two of the three original targets and conditions the third. The "
   "exposure target is met by design. The torque-accuracy target is achievable but, on its "
   "own, does not deliver the preload quality it appears to promise; what delivers it is the "
   "angle channel. The cycle-time target is not met by the architecture as first described, "
   "with one runner per flange, anywhere within the investigated parameter bounds. The %.1f h "
   "single-head result is not a failure of the concept but a finding about it: a single "
   "tool is the wrong architecture for a 92-bolt, multi-pass closure. The architecture the "
   "analysis supports has four synchronised heads per flange and one bolt-transfer channel "
   "per head, with the legacy tightening procedure unchanged, which brings bolting from %.1f h "
   "to %.1f h; a qualified two-star-pass procedure is an optional further step, not a "
   "requirement. Parallelisation, not faster tools, is what matters: the admissibility analysis "
   "shows that no plausible per-bolt time lets one head meet the target, and once torque "
   "application is parallelised every stage around it — bolt handling above all — "
   "must be parallelised too. This is a substantially larger and more expensive machine "
   "than the original proposal implied, and it changes the comparison with manual work: a "
   "single automated head is slower than a crew of four to six, and the automated system "
   "matches or improves on manual duration only when it too works several bolts at once. "
   "The target also has to be read as the source stated it, for bolting: the complete "
   "operation remains about %.1f h, governed by head handling and the leak test, and "
   "shortening it is a separate design problem for those subsystems."
   % (b1["bolting_total"]["median"], b1["bolting_total"]["median"], OB["bolting_total"]["median"],
      OE["full_total"]["median"]))
H2("6.2 Sequence choice as a throughput lever")
Pp("For a single head the star pattern costs far more travel than a circular pass (Fig. 2b), "
   "and circular-only tightening shortens the cycle, but the legacy procedure uses the star "
   "pattern for good reason: it limits the load redistribution caused by elastic "
   "interaction %s. Section 5.5 shows that optimising the order of the tool path under a "
   "load-spreading constraint recovers only a few per cent of the travel, because the "
   "opposite-bolt rule fixes most of it. Multi-head carriages remove the penalty while "
   "keeping cross-pattern loading, because each station tightens bolts at 90\u00b0 spacing "
   "together. The remaining sequence lever is the number of passes: with four heads and "
   "four transfer channels, removing one star pass shortens closing by %.0f%%. Alternative "
   "patterns with fewer passes %s and sequences optimised using measured interaction "
   "coefficients %s are the way to make that saving safely, and the manufacturing question "
   "becomes which reduced-pass procedure gives acceptable preload uniformity on this joint."
   % (C("bibel1992", "zhu2018"),
      100 * (1 - OC["bolting_close"]["median"] / OB["bolting_close"]["median"]),
      C("brown2010"), C("zhu2019")))
H2("6.3 Flange condition and engagement")
Pp("The engagement and alignment steps assume that nuts are where the rail expects them. "
   "Coke-drum bottom flanges can become distorted by thermal transients until bolt holes no "
   "longer align %s, and fouling of nuts by coke is likely. Both would raise the retry and "
   "operator-review probabilities, which the four-head sensitivity ranks among the stronger "
   "drivers. Vision-based nut localisation and pose estimation %s and self-adjusting "
   "engagement tools %s address part of this; a dimensional survey of the actual flanges "
   "before rail fabrication addresses the rest." % (C("horstmeyer2025"), C("yibang2023"),
                                                   C("ali2018")))
H2("6.4 Relation to unheading devices")
Pp("Where automatic bottom unheading devices are installed, the bottom head is no longer "
   "unbolted on every cycle %s, and the per-cycle bolting problem addressed here does not "
   "arise at that flange; the drum-to-device joint is then opened only at maintenance "
   "intervals, when flange distortion can make reassembly difficult %s. The architecture "
   "evaluated here is relevant to units that retain bolted heads, and its bolting subsystem "
   "could serve those maintenance reassemblies. Whether automated bolting or an unheading "
   "retrofit is preferable for a given unit is an economic and engineering decision outside "
   "the scope of this paper." % (C("berry2019"), C("horstmeyer2025")))
H2("6.5 Hazardous-area compliance and functional safety")
Pp("The source proposal described hazardous-area certification of the integrated system as "
   "achievable by selecting certified components. Component certification is necessary but "
   "not sufficient: the integrated assembly, including carriage drives, transducers, the "
   "carousel and the controller, must satisfy the applicable protection concepts as a whole "
   "%s. The gas-clearance and head-motion interlocks are safety functions, and their "
   "integrity requirements belong to a functional-safety assessment %s. The exhaustive check "
   "in Section 5.9 shows that the logic is consistent and deadlock-free as specified; it does "
   "not establish the hardware fault tolerance or safety integrity level of its "
   "implementation." % (C("iec60079", "atex", "singh2023"), C("iec61511")))
H2("6.6 Limitations")
Pp("Every performance figure in this paper is a model prediction. The step-time ranges in "
   "Table 2 are assumptions, bounded but not measured; the bolt-circle diameters are "
   "bracketed by a spacing rule; head masses are swept. The distributions and their "
   "independence are also assumptions; Section 5.6 shows that they move the tails of the "
   "cycle-time distribution but not the ranking of architectures. Elastic interaction is "
   "represented "
   "only through the number of passes, not modelled directly. The preload model treats nut "
   "factor and angle-estimate errors as independent and normally distributed. The FMEA "
   "scores are judgement. The control-logic verification abstracts the controller to a "
   "finite-state model and checks it for flanges of 4, 8 and 12 bolts; its conclusions "
   "concern the logic, not an implementation. The manual baseline of 4\u20136 h is as "
   "reported, not measured. Costs are not analysed beyond the preliminary estimates noted in "
   "Section 8.")

# ==================================================================== 7
H("7 Experimental validation and falsification plan")
Pp("The predictions above can be falsified by a staged test programme. Table 10 lists the "
   "measurements, the rig on which each is made and the acceptance criterion that would "
   "confirm or refute the corresponding result.")
TBL(["Quantity", "Rig", "Acceptance criterion"],
    [["Step times: align, engage, torque, release, verify, transfer",
      "Full-size single-station rig with a representative M56\u2013M64 stud and flange segment",
      "Measured medians within the Table 2 ranges; if not, re-run the model with measured values"],
     ["Carriage speed and indexing overhead", "Rail segment, loaded carriage",
      "As above"],
     ["Per-bolt-pass time, four-head station", "Four-head carriage on a full flange mock-up",
      "\u2264 %.0f s median to meet 2.5 h (Section 5.3)" % (4 * adm["budget_s_per_bolt_pass_parallel_2.5h"])],
     ["Bolt transfer with four channels", "Carousel with four access stations on the mock-up",
      "Median transfer time within the Table 2 range per channel, with no interference "
      "between channels (configuration B)"],
     ["Two-star-pass procedure and optimised tool path", "Full-flange mock-up with "
      "instrumented studs and the service gasket",
      "Final preload scatter and gasket-stress uniformity no worse than the legacy three-pass "
      "procedure; otherwise configuration B, not C–E, is the design"],
     ["Nut-factor scatter CV_K", "Instrumented studs (ultrasonic elongation) on the mock-up",
      "Measured CV_K; preload prediction of Fig. 8a to be confirmed within \u00b12 points"],
     ["Angle-estimate uncertainty CV_\u03b8", "As above, with gasket installed",
      "CV_\u03b8 \u2264 5% for the torque\u2013angle benefit of Fig. 8b to hold"],
     ["Engagement success on distorted flange", "Mock-up with imposed ovality and cupping",
      "Retry probability within Table 2 range"],
     ["Interlock logic", "Controller hardware-in-the-loop",
      "All seven invariants and the resume rule demonstrated; no deadlock after ESD in any phase"],
     ["Head handling", "Counterbalanced manipulator with dummy head at maximum mass",
      "Lift, swing, lock and reseat without personnel in load path; lock holds on air loss"]],
    widths=[1.8, 2.2, 2.3])
CAP("Table 10. Validation measurements and acceptance criteria. A pilot on one bottom head is "
    "proposed only after the single-station and mock-up tests have been passed.")

# ==================================================================== 8
H("8 Industrial feasibility")
Pp("Each subsystem uses established technology \u2014 pneumatic torque multipliers, rails "
   "and carriages, fastener magazines, counterbalanced manipulators and programmable "
   "controllers \u2014 so the development risk lies in integration and in the hazardous-area "
   "environment rather than in any single component. The preliminary cost estimates of the "
   "source proposal, about US$0.14\u20130.20 million for a single-flange pilot and "
   "US$0.74\u20131.1 million for all three flanges on four drums, assumed one runner per "
   "flange. The results of Section 5 imply a four-head carriage and four bolt-transfer "
   "channels on each flange if the bolting target is to be met, so those estimates are likely "
   "to be low and should be revised after the single-station tests; shortening the complete "
   "operation further would require faster head handling and leak testing, which were not "
   "costed. Rather than substitute another point estimate, the revised cost per flange is "
   "structured as")
EQ("C = C_fixed + r\u00b7C_head + c\u00b7C_channel + C_controls + C_certification")
Pp("with r = 4 heads and c = 4 transfer channels in the recommended configuration; the "
   "coefficients \u2014 rail and structure, each torque head, each transfer channel, the "
   "controls and the hazardous-area certification of the integrated assembly \u2014 require "
   "vendor quotations and are not estimated here. A staged programme \u2014 single-station "
   "rig, full-flange mock-up, then a pilot on one bottom head \u2014 allows the project to "
   "stop cheaply if the measured step times, nut-factor scatter or engagement success fall "
   "outside the ranges on which the results depend.", indent=False)

# ==================================================================== 9
H("9 Conclusions")
Pp("An automated architecture for opening and closing a 92-bolt, three-flange pressure "
   "vessel was evaluated by simulation against the targets originally claimed for it. With "
   "one torque head per flange, bolting takes a median %.1f h and cannot meet a 2.5 h target "
   "within the investigated parameter bounds: a single tool is the wrong architecture for "
   "this closure. "
   "Parallelising in stages brings bolting to %.1f h with four synchronised heads and %.1f h "
   "with a bolt-transfer channel per head. That machine-only configuration, with the legacy "
   "tightening procedure unchanged, is the recommended design: it meets the bolting target in "
   "%.0f%% of samples, and the ranking of architectures that leads to it holds in every "
   "sample under uniform, triangular and correlated input models. One fewer star pass would "
   "bring bolting to %.1f h but is an optional process change that requires qualification of "
   "preload uniformity. An exactly optimised tool path and pipelined servicing add little, the "
   "first because the opposite-bolt rule fixes most of the travel and the second because the "
   "bottom head is on the critical path in every phase. The complete operation remains about "
   "%.1f h, governed by head handling and the leak test, so the 2.5 h target holds for "
   "bolting but not for the whole operation; Eq. (%d) gives the servicing-time envelope a "
   "faster design would have to meet. Specifying \u00b12%% torque accuracy does little for "
   "preload, which nut-factor scatter dominates, whereas torque\u2013angle verification with "
   "re-torque substantially narrows it, to an extent that depends on how well joint stiffness "
   "is known. "
   "Explicit-state verification found no violation of the seven safety invariants but exposed "
   "widespread deadlock in the initial interlock specification after interrupted sequences; "
   "adding an explicit resume rule eliminated every deadlock while preserving all seven "
   "invariants. Pneumatic head handling requires a counterbalance carrying most of "
   "the load. The reduction in hazard-zone exposure, from %d\u2013%d person-hours to "
   "essentially none, is the most robust benefit. These are predictions, and the validation "
   "plan specifies the measurements that would confirm or overturn them."
   % (b1["bolting_total"]["median"], OA["bolting_total"]["median"], OB["bolting_total"]["median"],
      100 * OB["p_bolting_le_2.5h"], OC["bolting_total"]["median"], OE["full_total"]["median"],
      EQ_SERVICE, ex["manual_person_hours"][0], ex["manual_person_hours"][1]))

# ==================================================================== DECLARATIONS
H("Declarations")
for head, body in [
    ("Funding", "This work received no external funding."),
    ("Competing interests", "The author declares no competing interests."),
    ("Ethics approval", "Not applicable. No human participants or animals were involved."),
    ("Consent to participate / for publication", "Not applicable."),
    ("Prior dissemination", "An earlier, non-quantitative description of the architecture was "
     "submitted by the author in response to an industrial open-innovation challenge. That "
     "submission was not peer reviewed or published, and contained none of the models or "
     "results reported here."),
    ("Data availability", "No experimental data were generated. All model inputs are stated in "
     "the paper."),
    ("Code availability", "The models, the verification script and the figure generators are "
     "openly available at %s%s. Running pabhs_model.py, plc_verify.py and make_figures.py "
     "reproduces every number and figure."
     % (REPO_URL, (" and archived at https://doi.org/%s" % CODE_DOI) if CODE_DOI else "")),
    ("Author contributions", "L.S. conceived the study, developed the models, performed the "
     "analysis and wrote the manuscript."),
    ("Use of generative AI", "During the preparation of this work the author used Claude "
     "(Anthropic) to assist with literature search, implementation of the models and drafting "
     "of the text. The author reviewed and edited all content, verified every cited reference "
     "against its source, and takes full responsibility for the content of the article.")]:
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.3
    r = p.add_run(head + ". "); r.bold = True
    p.add_run(body)

# ==================================================================== REFERENCES
H("References")


def fmt(tag):
    if tag in STANDARDS:
        return STANDARDS[tag]
    v = DOIREFS[tag]
    au = v["authors"]
    aus = ", ".join(au[:6]) + (", et al." if len(au) > 6 else "")
    if not aus:
        aus = v.get("publisher") or ""
    title = v["title"].rstrip(".")
    kind = v.get("type")
    if kind in ("book", "monograph", "edited-book"):
        return "%s (%s) %s. %s. https://doi.org/%s" % (aus, v["year"], title, v["publisher"], v["doi"])
    if kind == "proceedings-article":
        ev = v.get("event") or ""
        if "SAE Technical Paper" in v["container"]:
            return "%s (%s) %s. SAE Technical Paper %s. %s. https://doi.org/%s" % (
                aus, v["year"], title, v["doi"].split("/")[-1], v["publisher"], v["doi"])
        cont = v["container"]
        if not ev or ev[:25].lower() in cont.lower() or cont[:25].lower() in ev.lower():
            where = cont                                   # container already names the event
        else:
            where = "Proceedings of the %s, %s" % (ev, cont)
        return "%s (%s) %s. In: %s. %s. https://doi.org/%s" % (aus, v["year"], title, where,
                                                               v["publisher"], v["doi"])
    loc = v["container"]
    if v.get("volume"):
        loc += " " + str(v["volume"])
        if v.get("issue"):
            loc += "(%s)" % v["issue"]
    if v.get("page"):
        loc += ":" + str(v["page"])
    return "%s (%s) %s. %s. https://doi.org/%s" % (aus, v["year"], title, loc, v["doi"])


for i, tag in enumerate(ORDER, 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    p.paragraph_format.space_after = Pt(2)
    p.add_run("%d. %s" % (i, fmt(tag))).font.size = Pt(9.5)

os.makedirs(OUT, exist_ok=True)
doc.core_properties.author = "Leon Sandler"
doc.core_properties.title = TITLE
path = os.path.join(OUT, "PABHS_IJAMT_Manuscript.docx")
doc.save(path)
words = sum(len(p.text.split()) for p in doc.paragraphs)
print("abstract words:", len(ABSTRACT.split()))
print("total words   :", words)
print("equations     :", EQN[0])
print("references    : %d cited; unused: %s" % (len(ORDER), sorted(set(DOIREFS) - set(ORDER))))
print("saved         :", path)
