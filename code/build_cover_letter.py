# -*- coding: utf-8 -*-
"""Cover letter for the IJAMT submission; numbers from the model output."""
import io
import json
import os

from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\PABHS_IJAMT"
R = json.load(io.open(os.path.join(HERE, "pabhs_results.json"), encoding="utf-8"))
V = json.load(io.open(os.path.join(HERE, "plc_verify_results.json"), encoding="utf-8"))
c = R["cycle"]
b1, b4 = c["1 head, parallel flanges"], c["4 heads, parallel flanges"]
ta = {(t["cv_K"], t["cv_theta"]): t for t in R["preload"]["torque_angle"]}

TITLE = ("Automated bolting and head handling for multi-flange pressure-vessel closures: "
         "a simulation-based design evaluation with a delayed-coking case study")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(9); st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1.0)


def P(t, bold=False, after=9):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(t); r.bold = bold


for line in ("Leon Sandler", "Independent Researcher", "Northbrook, Illinois, USA",
             "sandler.leon@gmail.com | ORCID 0009-0007-4584-808X"):
    P(line, after=0)
P("", after=4)
P("The Editor-in-Chief", after=0)
P("The International Journal of Advanced Manufacturing Technology", after=12)
P("Dear Editor,")
P("I am submitting the manuscript \u201c%s\u201d for consideration as an original research "
  "article." % TITLE)
P("Large bolted closures on process vessels are assembled and disassembled by hand in many "
  "plants, following multi-pass tightening procedures that make the operation a recurring "
  "manufacturing task in hazardous conditions. Automating it is often proposed, but rarely "
  "evaluated quantitatively. The manuscript takes one proposed architecture \u2014 orbital "
  "pneumatic torque runners, a bolt carousel, a counterbalanced head manipulator and an "
  "interlocking sequence controller \u2014 and tests each of its claimed targets by "
  "simulation for a 92-bolt, three-flange delayed-coking drum.")
P("Its main findings are:", after=4)
for t in [
    "A Monte Carlo cycle-time model built on the actual star-pattern rail travel shows that "
    "one torque head per flange needs a median %.1f h and cannot meet the 2.5 h target even "
    "with every step at its fastest assumed value; four synchronised heads per flange give "
    "%.1f h, and the carousel then becomes the rate-limiting element."
    % (b1["bolting_total"]["median"], b4["bolting_total"]["median"]),
    "Specifying \u00b12%% torque accuracy has almost no effect on preload scatter, which "
    "nut-factor variation dominates; torque\u2013angle verification with re-torque reduces the "
    "share of bolts outside \u00b110%% of target from %.0f%% to %.0f%%."
    % (100 * ta[(0.1, 0.03)]["outside10_before"], 100 * ta[(0.1, 0.03)]["outside10_after"]),
    "Exhaustive verification of the interlock logic finds no safety violation but reveals that "
    "the interlocks as first specified deadlock after an emergency stop mid-sequence (%d of %d "
    "reachable states for a 12-bolt flange); an explicit resume rule removes every deadlock, "
    "and mutation tests confirm that each invariant can fail."
    % (V["as_specified"]["12"]["deadlocked_states"], V["as_specified"]["12"]["states"]),
    "Pneumatic handling of a multi-tonne head is feasible only when a passive counterbalance "
    "carries most of the load."]:
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4)
    p.add_run(t)
P("I believe the work suits the journal because it addresses automated assembly and "
  "fastening, manufacturing-system simulation and control-logic verification, topics on "
  "which the journal has published, including robotic nut tightening and automated bolt "
  "engagement. The paper states plainly that its results are model predictions from "
  "stated assumptions, and it closes with a validation plan whose acceptance criteria could "
  "overturn them.")
P("An earlier, non-quantitative description of the architecture was submitted by me in "
  "response to an industrial open-innovation challenge. That submission was not peer "
  "reviewed or published and contained none of the models or results reported here. The "
  "manuscript is original, is not under consideration elsewhere, and has not been published. "
  "I am the sole author, I have no competing interests, and the work received no external "
  "funding. The models, verification script and figure generators are openly available at "
  "https://github.com/sandlerleon/pabhs-bolting-model. Generative AI (Claude, Anthropic) "
  "assisted with literature search, model implementation and drafting; I reviewed and "
  "edited all content, verified every reference against its source, and take full "
  "responsibility for the manuscript, as stated in its Declarations.")
P("Thank you for your consideration.", after=12)
P("Sincerely,", after=0)
P("Leon Sandler")
os.makedirs(OUT, exist_ok=True)
doc.core_properties.author = "Leon Sandler"
path = os.path.join(OUT, "PABHS_IJAMT_Cover_Letter.docx")
doc.save(path)
print("saved", path, "| words", sum(len(p.text.split()) for p in doc.paragraphs))
