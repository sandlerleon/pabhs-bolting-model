# -*- coding: utf-8 -*-
"""Check the built manuscript against the model outputs and against itself."""
import io
import json
import re

from docx import Document

DOC = r"C:\Users\Leon\Downloads\PABHS_IJAMT\PABHS_IJAMT_Manuscript.docx"
R = json.load(io.open("pabhs_results.json", encoding="utf-8"))
V = json.load(io.open("plc_verify_results.json", encoding="utf-8"))
d = Document(DOC)
paras = [p.text for p in d.paragraphs]
cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
T = "\n".join(paras + cells)
c = R["cycle"]
CHECK = [
    ("1-head median", "%.2f" % c["1 head, parallel flanges"]["bolting_total"]["median"]),
    ("4-head median", "%.2f" % c["4 heads, parallel flanges"]["bolting_total"]["median"]),
    ("config B bolting", "%.2f" % R["optimisation"]["B + 4 transfer channels"]["bolting_total"]["median"]),
    ("config C bolting", "%.2f" % R["optimisation"]["C + 2 star passes"]["bolting_total"]["median"]),
    ("config E complete", "%.2f" % R["optimisation"]["E + pipelined servicing"]["full_total"]["median"]),
    ("1-head best case", "%.2f" % R["best_case_bolting_h"]["1 head, parallel flanges"]),
    ("budget s/bolt-pass", "%.0f" % R["admissibility"]["budget_s_per_bolt_pass_parallel_2.5h"]),
    ("modelled s/bolt-pass", "%.0f" % R["admissibility"]["modelled_median_s_per_bolt_pass_bottom"]),
    ("carousel share", "%.0f%%" % (100 * R["carousel_share_4heads"])),
    ("deadlocked 12-bolt", str(V["as_specified"]["12"]["deadlocked_states"])),
    ("states 12-bolt", str(V["full_design"]["12"]["states"])),
    ("FMEA top RPN", str(R["fmea"][0]["RPN"])),
    ("coverage", "%.2f" % R["exposure"]["coverage"]),
]
bad = 0
for name, val in CHECK:
    ok = val in T
    bad += not ok
    print("  %-22s %-8s %s" % (name, val, "ok" if ok else "NOT FOUND"))
i = paras.index("References")
body = "\n".join(paras[:i] + cells)
refs = [p for p in paras[i + 1:] if re.match(r"^\d+\. ", p)]
cited = set()
for m in re.finditer(r"\[([\d,\s\u2013-]+)\]", body):
    for part in m.group(1).split(","):
        part = part.strip()
        if "\u2013" in part:
            a, b = part.split("\u2013")
            cited |= set(range(int(a), int(b) + 1))
        elif part.isdigit():
            cited.add(int(part))
listed = set(range(1, len(refs) + 1))
print("  references: %d listed, never cited %s, cited-not-listed %s"
      % (len(refs), sorted(listed - cited) or "none", sorted(cited - listed) or "none"))
figs = set(int(x) for x in re.findall(r"Fig\. (\d+)", body))
tabs = set(int(x) for x in re.findall(r"Table (\d+)", body))
n_fig = len({int(m.group(1)) for p_ in paras for m in [re.match(r"^Fig\. (\d+) ", p_)] if m})
n_tab = len({int(m.group(1)) for p_ in paras for m in [re.match(r"^Table (\d+)\. ", p_)] if m})
if figs != set(range(1, n_fig + 1)) or tabs != set(range(1, n_tab + 1)):
    print("  !! figure/table citations do not match captions:", n_fig, n_tab)
    bad += 1
first = []
for m in re.finditer(r"Fig\. (\d+)", body):
    if int(m.group(1)) not in first:
        first.append(int(m.group(1)))
if first != sorted(first):
    print("  !! figures first cited out of order:", first)
    bad += 1
print("  figures cited:", sorted(figs), " tables cited:", sorted(tabs))
for phrase in ("guaranteed", "eliminates manual personnel exposure", "TRL 7", "Wazoku",
               "InnoCentive", "YPF", "winning"):
    if phrase.lower() in T.lower():
        print("  !! phrase from the source proposal survives:", phrase)
        bad += 1
for p in paras:
    if re.search(r"%[sd]|%\.\d|%%", p):
        print("  !! unfilled format placeholder:", p[:90])
        bad += 1
print("VERDICT:", "PASS" if not bad and listed == cited else "REVIEW")
