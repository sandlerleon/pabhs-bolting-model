# -*- coding: utf-8 -*-
"""Repair the Rubriq-edited IJAMT manuscript (PABHS_IJAMT_Manuscript_v2.docx) without
discarding its style edits (US spelling, commas, articles, 'multipass').

Rubriq's pass also changed the sense of some thirty passages, broke several sentences,
changed a term of art ('explicit-state model checking'), made the sole author plural,
reversed a cost statement ('not costed' -> 'not costly'), and altered three reference
titles. Each is restored below; every fix must hit at least once.

    python fix_rubriq_v2.py
"""
import os
import shutil

from docx import Document

D = r"C:\Users\Leon\Downloads\PABHS_IJAMT"
RAW = os.path.join(D, "_rubriq_raw", "PABHS_IJAMT_Manuscript_v2_rubriq_raw.docx")
OUT = os.path.join(D, "PABHS_IJAMT_Manuscript_v2.docx")

FIXES = [
    # title: keep identical to the Zenodo record, cover letter and repository
    ("multiflange pressure-vessel closures: A simulation-based",
     "multi-flange pressure-vessel closures: a simulation-based"),
    ("a bolt-managed carousel", "a bolt-management carousel"),
    ("after interrupted sequences, for which a resume rule was removed.",
     "after interrupted sequences, which a resume rule removed."),
    # the unheading-device trend does not cause re-tightening
    ("automatic bottom unheading devices such that the joints that remain",
     "automatic bottom unheading devices, and the joints that remain"),
    # one operation every 12 h for the unit, not every drum every 12 h
    ("on a 48 h fill cycle are opened or closed every 12 h,",
     "on a 48 h fill cycle give an opening or closing operation every 12 h,"),
    ("challenge, such as a pneumatic", "challenge, as a pneumatic"),
    ("interaction coefficients allow a one-pass", "interaction coefficients allowed a one-pass"),
    # re-presentation = presenting the bolt again for reinstallation
    ("storage, representation, installation", "storage, re-presentation, installation"),
    ("with N bolts sit at angle", "with N bolts sits at angle"),
    ("(wrench clearance is set at the lower end and gasket load uniformity is set at the upper end)",
     "(wrench clearance sets the lower end, gasket-load uniformity the upper)"),
    ("red arrows indicate the first eight carriages, and the orange arrows indicate one group",
     "red arrows indicate the first eight carriage moves, and orange marks one group"),
    ("Opening is a one-stage breakout pass", "Opening is one staged breakout pass"),
    ("a circular run-on pass, and three stars pass at increasing torque",
     "a circular run-on pass, three star passes at increasing torque"),
    ("and the carousel is N\u00b7t_transfer/c", "and the carousel takes N\u00b7t_transfer/c"),
    ("each by their own runner", "each by its own runner"),
    ("With respect to all the inputs in Table 2, the spacing factor k and the position of each "
     "bolt within its range", "All inputs in Table 2, the spacing factor k and the position of "
     "each bolt size within its range"),
    ("[22] is followed by b = 320 mm", "[22] follows by setting b = 320 mm"),
    ("This is an explicit-state model for checking the logic as specified",
     "This is explicit-state model checking of the logic as specified"),
    ("When the flanges are in series nearly doubled to 14.5 h",
     "Working the flanges in series nearly doubles this to 14.5 h"),
    ("In terms of head handling, gasket", "Including head handling, gasket"),
    ("torque per pass (0.43)", "torque time per pass (0.43)"),
    ("Table 5. Remove the bottlenecks", "Table 5. Removing the bottlenecks"),
    ("In turn, the bottlenecks are removed from Table 5 and Fig. 6. Given each of the four heads",
     "Table 5 and Fig. 6 remove the bottlenecks in turn. Giving each of the four heads"),
    ("diameters and bolts from 7.96 to", "diameters and bolting from 7.96 to"),
    ("the residual is approximately 10%, with 8% of the benefit largely disappearing.",
     "the residual is approximately 10%; with 8%, the benefit largely disappears."),
    ("589 mm above the largest standard bore", "589 mm, beyond the largest standard bore"),
    ("The head mass was not visible and was swept.", "The head mass was not available and was swept."),
    ("The scores are the authors\u2019 judgments", "The scores are the author's judgement"),
    ("as the source stated it; for bolting, the complete operation remains at approximately 3.4 h",
     "as the source stated it, for bolting: the complete operation remains approximately 3.4 h"),
    ("to make that safe, and the manufacturing", "to make that saving safely, and the manufacturing"),
    ("Both increase the retry", "Both would raise the retry"),
    ("with the four-head sensitivity ranking among the stronger drivers",
     "which the four-head sensitivity ranks among the stronger drivers"),
    ("FMEA scores are judged.", "FMEA scores are the author's judgement."),
    ("which were not costly. Rather than substituting", "which were not costed. Rather than substituting"),
    ("imply four-head carriage", "imply a four-head carriage"),
    ("programmable controllers\u2014so that the development risk",
     "programmable controllers\u2014so the development risk"),
    ("One fewer start pass", "One fewer star pass"),
    ("a faster design that would have to meet", "that a faster design would have to meet"),
    # reference titles are quoted exactly as published
    ("Head Flange Design Optimization", "Head Flange Design Optimisation"),
    ("with multidiameter bolts/nuts", "with multi-diameter bolts/nuts"),
    ("on the harmonization of the laws", "on the harmonisation of the laws"),
]


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def replace_in_paragraph(p, old, new):
    runs = p.runs
    text = "".join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos, first = 0, None
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        s, e = max(i, a) - a, min(j, b) - a
        if first is None:
            r.text = r.text[:s] + new + r.text[e:]
            first = r
        else:
            r.text = r.text[:s] + r.text[e:]
    return True


def main():
    if not os.path.exists(RAW):
        os.makedirs(os.path.dirname(RAW), exist_ok=True)
        shutil.copy2(OUT, RAW)
    doc = Document(RAW)
    seen = set()
    for old, new in FIXES:
        hits = 0
        for p in all_paragraphs(doc):
            if id(p._p) in seen and False:
                continue
            while old in p.text:
                assert replace_in_paragraph(p, old, new)
                hits += 1
        assert hits >= 1, "no hit: %r" % old[:70]
        print("%2d  %s" % (hits, old[:70]))
    doc.core_properties.author = "Leon Sandler"
    doc.save(OUT)
    print("repaired:", OUT)


if __name__ == "__main__":
    main()
